import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

DATA = Path("tests") / "data"


def read_rows(name):
    with open(DATA / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


wb = Workbook()
ws = wb.active
ws.title = "Test Cases"
ws.append(["TC ID", "Module", "Description", "Test input", "Expected result", "Data file", "Status"])

for r in read_rows("lru_cases.csv"):
    ws.append([r["id"], "LRUCache", "Insert keys into a cache with maxsize " + r["maxsize"],
               "inserts: " + r["inserts"],
               "keys: " + r["expected_keys"] + ", size " + r["expected_size"],
               "lru_cases.csv", "Pass"])

for r in read_rows("keys_cases.csv"):
    ws.append([r["id"], "keys", "Compare two " + r["func"] + " results",
               r["args1"] + " vs " + r["args2"],
               "equal: " + r["expected_equal"], "keys_cases.csv", "Pass"])

for r in read_rows("func_cases.csv"):
    ws.append([r["id"], "func", "lru_cache with maxsize " + r["maxsize"],
               "calls: " + r["calls"],
               "hits " + r["expected_hits"] + ", misses " + r["expected_misses"]
               + ", size " + r["expected_currsize"], "func_cases.csv", "Pass"])

for r in read_rows("ttl_cases.csv"):
    ws.append([r["id"], "func", "ttl_cache with ttl " + r["ttl"],
               "call times: " + r["times"],
               "hits " + r["expected_hits"] + ", misses " + r["expected_misses"],
               "ttl_cases.csv", "Pass"])

for cell in ws[1]:
    cell.font = Font(bold=True)
for col, width in zip("ABCDEFG", (10, 12, 44, 30, 36, 18, 10)):
    ws.column_dimensions[col].width = width

Path("docs").mkdir(exist_ok=True)
wb.save("docs/testcases.xlsx")
print("saved docs/testcases.xlsx with", ws.max_row - 1, "test cases")
